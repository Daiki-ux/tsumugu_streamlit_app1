# docs/slides/merge_deck.py — 前段資料とシステムパートを 1 つの .pptx にまとめる
#   使い方: python merge_deck.py 前段.pptx システムパート.pptx 出力.pptx [先頭から移す枚数（既定 1、別案は 2）]
#   ・前段のデザイン（テーマ・レイアウト）を土台に、システムパートのスライドを後ろへ足す
#   ・システムパートの 1 枚目（エグゼクティブサマリ）は全体の 2 枚目へ移す
#   ・前段 10 枚目（導入によって期待される効果）はシステムパートの効果の 1 枚に統合したので外す
#   ・前段 9 枚目の「担当者が気づきを残す」を「日報に書かれた気づきが残る」に言い換える
#   ・ページ番号（右下の数字）を通しで振り直す
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

FRONT, SYSTEM, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
DROP_FRONT = [10]           # 外す前段のスライド番号（1 始まり）
SUMMARY_TO = 2              # システムパートの先頭（サマリ）を置く位置
LEAD = int(sys.argv[4]) if len(sys.argv) > 4 else 1   # 先頭から移す枚数

tmp = Path(tempfile.mkdtemp())
base, sysd = tmp / "base", tmp / "sys"
zipfile.ZipFile(FRONT).extractall(base)
zipfile.ZipFile(SYSTEM).extractall(sysd)


def read(p): return Path(p).read_text(encoding="utf-8")
def write(p, s): Path(p).parent.mkdir(parents=True, exist_ok=True); Path(p).write_text(s, encoding="utf-8")
def count(d, pat): return len(list(Path(d).glob(pat)))


ct = read(base / "[Content_Types].xml")
pres = read(base / "ppt/presentation.xml")
prels = read(base / "ppt/_rels/presentation.xml.rels")

n_slides = count(base / "ppt/slides", "slide*.xml")
n_notes = count(base / "ppt/notesSlides", "notesSlide*.xml")
n_charts = count(base / "ppt/charts", "chart*.xml")
n_emb = count(base / "ppt/embeddings", "*.xlsx")
sys_slides = sorted((sysd / "ppt/slides").glob("slide*.xml"), key=lambda p: int(re.findall(r"\d+", p.name)[0]))

next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', prels)) + 1
next_sid = max(int(x) for x in re.findall(r'<p:sldId id="(\d+)"', pres)) + 1
new_ids = []
SLIDE_CT = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
NOTES_CT = "application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml"
CHART_CT = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"

for sp in sys_slides:
    n_slides += 1
    new = f"slide{n_slides}.xml"
    shutil.copy(sp, base / "ppt/slides" / new)
    rels = read(sysd / "ppt/slides/_rels" / (sp.name + ".rels"))

    def fix(m):
        global n_notes, n_charts, n_emb
        typ, target = m.group(1), m.group(2)
        if typ.endswith("/slideLayout"):
            return m.group(0).replace(target, "../slideLayouts/slideLayout1.xml")
        if typ.endswith("/notesSlide"):
            n_notes += 1
            src = sysd / "ppt/notesSlides" / Path(target).name
            dst = f"notesSlide{n_notes}.xml"
            shutil.copy(src, base / "ppt/notesSlides" / dst)
            nr = read(sysd / "ppt/notesSlides/_rels" / (src.name + ".rels"))
            nr = re.sub(r'Target="\.\./notesMasters/[^"]+"', 'Target="../notesMasters/notesMaster1.xml"', nr)
            nr = re.sub(r'Target="\.\./slides/[^"]+"', f'Target="../slides/{new}"', nr)
            write(base / "ppt/notesSlides/_rels" / (dst + ".rels"), nr)
            globals()["ct_add"].append((f"/ppt/notesSlides/{dst}", NOTES_CT))
            return m.group(0).replace(target, f"../notesSlides/{dst}")
        if typ.endswith("/chart"):
            n_charts += 1
            src = sysd / "ppt/charts" / Path(target).name
            dst = f"chart{n_charts}.xml"
            shutil.copy(src, base / "ppt/charts" / dst)
            crp = sysd / "ppt/charts/_rels" / (src.name + ".rels")
            if crp.exists():
                cr = read(crp)
                for emb in re.findall(r'Target="\.\./embeddings/([^"]+)"', cr):
                    n_emb += 1
                    ename = f"tsumugu_sys_{n_emb}.xlsx"
                    shutil.copy(sysd / "ppt/embeddings" / emb, base / "ppt/embeddings" / ename)
                    cr = cr.replace(f"../embeddings/{emb}", f"../embeddings/{ename}")
                write(base / "ppt/charts/_rels" / (dst + ".rels"), cr)
            globals()["ct_add"].append((f"/ppt/charts/{dst}", CHART_CT))
            return m.group(0).replace(target, f"../charts/{dst}")
        raise SystemExit(f"想定外の関係: {typ} {target}")

    ct_add = [(f"/ppt/slides/{new}", SLIDE_CT)]
    rels = re.sub(r'<Relationship [^>]*Type="([^"]+)"[^>]*Target="([^"]+)"[^>]*/>', fix, rels)
    write(base / "ppt/slides/_rels" / (new + ".rels"), rels)
    for part, typ in ct_add:
        ct = ct.replace("</Types>", f'<Override PartName="{part}" ContentType="{typ}"/></Types>')
    rid = f"rId{next_rid}"; next_rid += 1
    prels = prels.replace("</Relationships>",
        f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/{new}"/></Relationships>')
    new_ids.append(f'<p:sldId id="{next_sid}" r:id="{rid}"/>'); next_sid += 1

# スライドの並びを組み直す
lst = re.search(r"<p:sldIdLst>(.*?)</p:sldIdLst>", pres, re.S)
front_ids = re.findall(r"<p:sldId [^>]*/>", lst.group(1))
for i, x in enumerate(front_ids, 1):    # 外すスライドは関係ごと消す（残すと保存時に名前が重なる）
    if i not in DROP_FRONT:
        continue
    rid = re.search(r'r:id="([^"]+)"', x).group(1)
    target = re.search(rf'<Relationship [^>]*Id="{rid}"[^>]*Target="([^"]+)"[^>]*/>', prels).group(1)
    prels = re.sub(rf'<Relationship [^>]*Id="{rid}"[^>]*/>', "", prels)
    slide = base / "ppt" / target
    srels = slide.parent / "_rels" / (slide.name + ".rels")
    for nt in re.findall(r'Target="\.\./notesSlides/([^"]+)"', read(srels)):
        (base / "ppt/notesSlides" / nt).unlink()
        (base / "ppt/notesSlides/_rels" / (nt + ".rels")).unlink()
        ct = ct.replace(f'<Override PartName="/ppt/notesSlides/{nt}" ContentType="{NOTES_CT}"/>', "")
    slide.unlink(); srels.unlink()
    ct = ct.replace(f'<Override PartName="/ppt/{target}" ContentType="{SLIDE_CT}"/>', "")
front_ids = [x for i, x in enumerate(front_ids, 1) if i not in DROP_FRONT]
order = front_ids[:SUMMARY_TO - 1] + new_ids[:LEAD] + front_ids[SUMMARY_TO - 1:] + new_ids[LEAD:]
pres = pres.replace(lst.group(0), "<p:sldIdLst>" + "".join(order) + "</p:sldIdLst>")

write(base / "[Content_Types].xml", ct)
write(base / "ppt/presentation.xml", pres)
write(base / "ppt/_rels/presentation.xml.rels", prels)

packed = tmp / "packed.pptx"
with zipfile.ZipFile(packed, "w", zipfile.ZIP_DEFLATED) as z:
    for f in sorted(base.rglob("*")):
        if f.is_file():
            z.write(f, f.relative_to(base).as_posix())

# 文言の修正とページ番号の振り直し（python-pptx）
from pptx import Presentation
from pptx.util import Emu

p = Presentation(packed)
REPLACE = {
    "担当者が気づきを残す": "日報に書かれた気づきが残る",
    "本プロジェクトは、担当者の気づきを記録し、次の世代・本部へとつなぐことで、組織としての顧客理解を積み上げていく基盤である。":
        "本プロジェクトは、いつもの日報に書かれた担当者の気づきを残し、次の世代・本部へとつなぐことで、組織としての顧客理解を積み上げていく基盤である。",
}
for i, slide in enumerate(p.slides, 1):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            for r in para.runs:
                if r.text in REPLACE:
                    r.text = REPLACE[r.text]
        txt = sh.text_frame.text.strip()
        near = abs(Emu(sh.left).inches - 11.87) < 0.1 and abs(Emu(sh.top).inches - 6.92) < 0.1
        if near and txt.isdigit():
            runs = [r for para in sh.text_frame.paragraphs for r in para.runs]
            runs[0].text = str(i)
            for r in runs[1:]:
                r.text = ""
p.save(OUT)
print("wrote", OUT, "slides:", len(p.slides))
